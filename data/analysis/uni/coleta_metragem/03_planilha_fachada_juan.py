# -*- coding: utf-8 -*-
"""Planilha de fachada para o Juan (2026-09-29): unidades maduras (e jovens com >= 6 meses e metragem) SEM nota de
fachada de Felipe/Vini, as cegas (sem faturamento/alunos), + 6 unidades ja avaliadas misturadas sem aviso para medir a
concordancia (gabarito em _gabarito_nao_abrir/). Mesmo checklist e nota da coleta de 28/09.
Saida: planilha_fachada_juan.xlsx. Rodar com PYTHONPATH=src (.venv).
"""
import importlib.util
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).resolve().parent
UNI = AQUI.parent
RNG = np.random.default_rng(20260929)
SN = ["sim", "não", "não dá para ver"]
ITENS = ["letreiro_visivel_da_rua", "vitrine_para_rua", "frente_para_via_principal", "esquina", "estacionamento_proprio", "acesso_pelo_terreo"]
PREENCHER = ITENS + ["nota_fachada_1a5", "data_imagem_rua", "observacao"]
INSTRUCOES = [
    ("Objetivo", "Avaliar a fachada das unidades pelo Street View. Vamos testar se fachada melhor explica unidade que fatura acima do que a praça prevê."),
    ("Regra 1 — às cegas", "Não consulte faturamento, alunos ou ranking das unidades enquanto preenche. A planilha não traz nada disso de propósito."),
    ("Regra 2 — sem trocar ideia", "Avalie sozinho, sem ver a avaliação de outra pessoa. Algumas unidades já foram avaliadas por outra pessoa, de propósito, para medir se as réguas batem."),
    ("Como abrir", "Clique em 'abrir' na coluna link_rua (Street View). Se não abrir na frente da academia, use o link_satelite e arraste o bonequinho para a rua da entrada."),
    ("Checklist (antes da nota)", "letreiro visível da rua? vitrine/vidro para a rua? frente para avenida ou via principal? fica na esquina? tem estacionamento próprio? a entrada é no térreo? Use 'não dá para ver' quando não houver imagem."),
    ("Nota de fachada", "Sua impressão geral de 1 (escondida, sem presença) a 5 (muito visível e atraente). Responda o checklist antes da nota. Use números inteiros."),
    ("Data da imagem", "O Street View mostra a data da imagem (canto). Se for anterior à inauguração da academia, anote e use 'não dá para ver' no que não aparecer."),
    ("Tempo", "~2 minutos por unidade. Comece pela P1 (maduras), depois a P2."),
]


def main():
    s = importlib.util.spec_from_file_location("ex", UNI / "uni01" / "01_exploracao.py"); ex = importlib.util.module_from_spec(s); s.loader.exec_module(ex)
    _argv = sys.argv; sys.argv = [_argv[0], "2026-09-29"]
    s2 = importlib.util.spec_from_file_location("cons", AQUI / "02_consolidar.py"); cons = importlib.util.module_from_spec(s2); s2.loader.exec_module(cons)
    sys.argv = _argv
    f, d = ex.carregar()
    t = cons.ler()
    ja = set(t.loc[t.nota_fachada_1a5.notna(), "unidade_id"])
    mad = set(ex.desfechos(f, "2025-01", "2025-12").index) | set(ex.desfechos(f, "2026-01", "2026-08").index)
    est = pd.read_csv(AQUI / "metragem_estimada_2026-09-29.csv", sep=";", encoding="utf-8-sig", index_col=0).metragem_estimada_m2.dropna()
    tem_m2 = set(d.tam_metragem_m2.dropna().index) | set(est.index)
    mvmax = f[f.valido].groupby("unidade_id").mv.max()
    jovem6 = {u for u in d.index if u not in mad and mvmax.get(u, 0) >= 6 and u in tem_m2}
    ok = d.lat.notna() & d.lng.notna()
    p1 = sorted(u for u in mad if u not in ja and ok.get(u, False))
    p2 = sorted(u for u in jovem6 if u not in ja and ok.get(u, False))
    cal_pool = sorted(u for u in ja if u in mad and u != "botafogo-rj")
    cal = list(RNG.choice(cal_pool, size=min(6, len(cal_pool)), replace=False))
    base = pd.DataFrame({"unidade_id": p1 + p2 + cal, "prioridade": ["P1"] * len(p1) + ["P2"] * len(p2) + ["P1"] * len(cal)})
    base = base.sample(frac=1, random_state=20260929).sort_values("prioridade", kind="stable").reset_index(drop=True)
    dd = d.reindex(base.unidade_id)
    base["nome"], base["cidade"], base["uf"] = dd.nome.values, dd.cidade_cadastro.values, dd.uf.values
    base["inaugurou"] = pd.to_datetime(dd.inauguracao.values).strftime("%Y-%m")
    base["em_shopping_pista"] = np.where(dd.poi_em_shopping.map({True: True, False: False}).fillna(False).values, "provável shopping", "")
    base["link_rua"] = [f"https://www.google.com/maps/@?api=1&map_action=pano&viewpoint={la:.6f},{lo:.6f}" for la, lo in zip(dd.lat, dd.lng)]
    base["link_satelite"] = [f"https://www.google.com/maps/@{la:.6f},{lo:.6f},90m/data=!3m1!1e3" for la, lo in zip(dd.lat, dd.lng)]
    for c in PREENCHER:
        base[c] = None
    (AQUI / "_gabarito_nao_abrir").mkdir(exist_ok=True)
    pd.DataFrame({"unidade_id": cal}).to_csv(AQUI / "_gabarito_nao_abrir" / "fachada_juan_calibracao.csv", sep=";", encoding="utf-8-sig", index=False)
    arq = AQUI / "planilha_fachada_juan.xlsx"
    with pd.ExcelWriter(arq, engine="openpyxl") as w:
        pd.DataFrame(INSTRUCOES, columns=["tema", "como fazer"]).to_excel(w, sheet_name="Instrucoes", index=False)
        base.to_excel(w, sheet_name="Juan", index=False)
    wb = load_workbook(arq)
    wi = wb["Instrucoes"]; wi.column_dimensions["A"].width = 26; wi.column_dimensions["B"].width = 130
    ws = wb["Juan"]; hdr = {c.value: c.column_letter for c in ws[1]}; n = ws.max_row
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = Font(bold=True)
        if c.value in PREENCHER:
            c.fill = amarelo
    for col, w_ in (("unidade_id", 24), ("nome", 26), ("cidade", 18), ("link_rua", 10), ("link_satelite", 12), ("observacao", 40), ("em_shopping_pista", 16)):
        ws.column_dimensions[hdr[col]].width = w_
    for i in range(2, n + 1):
        for col in ("link_rua", "link_satelite"):
            cel = ws[f"{hdr[col]}{i}"]; cel.hyperlink = cel.value; cel.value = "abrir"; cel.font = Font(color="0563C1", underline="single")
    for campo in ITENS:
        dv = DataValidation(type="list", formula1='"' + ",".join(SN) + '"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"{hdr[campo]}2:{hdr[campo]}{n}")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True); ws.add_data_validation(dv)
    dv.add(f"{hdr['nota_fachada_1a5']}2:{hdr['nota_fachada_1a5']}{n}")
    ws.freeze_panes = "C2"
    wb.save(arq)
    print(f"P1 (maduras sem nota): {len(p1)} | P2 (jovens >= 6 meses com metragem): {len(p2)} | controle de concordancia: {len(cal)} | total {len(base)}")
    print(base[["prioridade", "nome", "uf", "inaugurou"]].to_string(index=False))


if __name__ == "__main__":
    main()
