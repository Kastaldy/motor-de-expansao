# -*- coding: utf-8 -*-
"""Gera a planilha de coleta manual de metragem (satelite), tipo de imovel e fachada — Felipe e Vini.

Desenho:
- Unidades SEM metragem, por prioridade: P1 = inauguradas ate ago/2025 (confirmam o UNI-02 em 2026); P2 = set/2025 a
  mar/2026; P3 = abr/2026 em diante.
- CALIBRACAO escondida: 12 unidades COM metragem conhecida (10 de rua em tamanhos variados + 2 em shopping), sem
  marcacao na planilha. O valor fica so' em gabarito_calibracao.csv (nao entregar aos avaliadores).
- SOBREPOSICAO: 10 unidades medidas pelos dois (6 de calibracao + 4 da P1), para medir a consistencia entre avaliadores.
- CEGO: a planilha nao tem faturamento nem nada de desempenho; a ordem e' embaralhada.
Saida: planilha_coleta_metragem.xlsx (abas Instrucoes, Felipe, Vini), gabarito_calibracao.csv, atribuicao.csv.
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).resolve().parent
UNI = AQUI.parent
RNG = np.random.default_rng(20260928)

TIPOS = ["loja de rua (térreo)", "galeria / centro comercial", "dentro de shopping", "andar superior de prédio",
         "anexo a supermercado/atacarejo", "prédio/galpão próprio", "outro"]
SN = ["sim", "não", "não dá para ver"]
CONF = ["alta", "média", "baixa"]
PREENCHER = ["mede_por_satelite", "area_telhado_m2", "pavimentos_academia", "metragem_estimada_m2", "confianca_medida",
             "tipo_imovel", "letreiro_visivel_da_rua", "vitrine_para_rua", "frente_para_via_principal", "esquina",
             "estacionamento_proprio", "acesso_pelo_terreo", "nota_fachada_1a5", "data_imagem_satelite", "observacao"]

INSTRUCOES = [
    ("Objetivo", "Estimar a metragem das unidades que não têm esse dado, e anotar tipo de imóvel e fachada. É o dado que falta para confirmar o modelo de faturamento (UNI-02)."),
    ("Regra 1 — às cegas", "Não consulte o faturamento, os alunos ou o ranking das unidades enquanto preenche. A planilha não traz nada disso de propósito."),
    ("Regra 2 — não troque linhas", "Cada um preenche só a própria aba. Algumas unidades aparecem nas duas abas de propósito: meça sem consultar o outro."),
    ("Regra 3 — sem pesquisar a metragem", "Não procure a metragem em site, anúncio ou planilha. O teste depende de ser medida pelo satélite."),
    ("Área do telhado", "Abra o link 'satélite' ou o Google Earth Web (earth.google.com). Use a régua → Polígono e contorne só o telhado da academia (sem estacionamento, calçada ou a loja vizinha). Anote a área em m²."),
    ("Pavimentos", "Pelo Street View (link 'rua'), conte quantos andares a ACADEMIA ocupa, não o prédio. Mezanino conta como 0,5."),
    ("Metragem estimada", "Calculada sozinha: área do telhado × pavimentos."),
    ("Quando não dá para medir", "Dentro de shopping, em andar de prédio com outras empresas ou sem imagem clara: marque 'mede_por_satelite' = não, deixe a área vazia e preencha o resto (tipo de imóvel, fachada)."),
    ("Confiança", "alta = contorno claro e só da academia; média = telhado compartilhado ou borda incerta; baixa = chute."),
    ("Tipo de imóvel", "Escolha na lista. Na dúvida, 'outro' e explique na observação."),
    ("Fachada (checklist)", "Responda pelo Street View: letreiro visível da rua? vitrine/vidro para a rua? frente para avenida ou via principal? fica na esquina? tem estacionamento próprio? a entrada é no térreo? Use 'não dá para ver' quando não houver imagem."),
    ("Nota de fachada", "Opcional: sua impressão geral de 1 (escondida) a 5 (muito visível e atraente). Responda o checklist antes da nota."),
    ("Data da imagem", "No Google Earth/Street View aparece a data da imagem. Se for anterior à inauguração, anote e marque confiança baixa."),
    ("Prioridade", "Comece pela P1 (as que confirmam o modelo em 2026), depois P2 e P3. Tempo estimado: ~5 min por unidade."),
]


def main():
    d = pd.read_parquet(UNI / "dim_unidade.parquet")
    d = d[~d.excluida_rede & d.lat.notna()].set_index("unidade_id")
    sem = d[d.tam_metragem_m2.isna()].copy()
    ina = sem.inauguracao
    sem["prioridade"] = np.where(ina <= "2025-08-31", "P1", np.where(ina <= "2026-03-31", "P2", "P3"))
    com = d[d.tam_metragem_m2.notna()].copy()
    # calibracao: 10 de rua espalhadas por tamanho + 2 em shopping
    rua = com[com.poi_em_shopping.fillna(0) == 0].sort_values("tam_metragem_m2")
    idx = np.unique(np.linspace(0, len(rua) - 1, 10).round().astype(int))
    cal = list(rua.index[idx]) + list(RNG.choice(com[com.poi_em_shopping.fillna(0) == 1].index, 2, replace=False))
    base = pd.concat([sem.assign(grupo="medir"), com.loc[cal].assign(grupo="calibracao", prioridade="P1")])
    # sobreposicao e divisao
    p1 = [u for u in sem.index if sem.at[u, "prioridade"] == "P1"]
    ambos = list(RNG.choice(cal, 6, replace=False)) + list(RNG.choice(p1, min(4, len(p1)), replace=False))
    resto = [u for u in base.index if u not in ambos]
    resto = sorted(resto, key=lambda u: (base.at[u, "prioridade"], RNG.random()))
    aval = {u: ("Felipe" if i % 2 == 0 else "Vini") for i, u in enumerate(resto)}
    aval.update({u: "ambos" for u in ambos})
    base["avaliador"] = pd.Series(aval)
    base["em_shopping_pista"] = np.where(base.poi_em_shopping.fillna(0) == 1, "provável (verificar)", "")
    base["link_satelite"] = [f"https://www.google.com/maps/@{la:.6f},{lo:.6f},90m/data=!3m1!1e3" for la, lo in zip(base.lat, base.lng)]
    base["link_rua"] = [f"https://www.google.com/maps/@?api=1&map_action=pano&viewpoint={la:.6f},{lo:.6f}" for la, lo in zip(base.lat, base.lng)]
    base["inaugurou"] = base.inauguracao.dt.strftime("%m/%Y")
    base["cidade"] = base.cidade_cadastro.fillna("")
    base[["grupo", "avaliador", "prioridade"]].to_csv(AQUI / "_gabarito_nao_abrir" / "atribuicao.csv", sep=";", encoding="utf-8-sig")
    com.loc[cal, ["nome", "tam_metragem_m2", "poi_em_shopping"]].to_csv(AQUI / "_gabarito_nao_abrir" / "gabarito_calibracao.csv", sep=";", encoding="utf-8-sig")

    arq = AQUI / "planilha_coleta_metragem.xlsx"
    cols = ["prioridade", "nome", "cidade", "uf", "inaugurou", "em_shopping_pista", "link_satelite", "link_rua"]
    with pd.ExcelWriter(arq, engine="openpyxl") as w:
        pd.DataFrame(INSTRUCOES, columns=["tema", "como fazer"]).to_excel(w, sheet_name="Instrucoes", index=False)
        for quem in ("Felipe", "Vini"):
            t = base[base.avaliador.isin([quem, "ambos"])].copy()
            t = t.assign(_o=[RNG.random() for _ in range(len(t))]).sort_values(["prioridade", "_o"])
            out = t[cols].reset_index().rename(columns={"index": "unidade_id"})
            for c in PREENCHER:
                out[c] = None
            out.to_excel(w, sheet_name=quem, index=False)
    wb = load_workbook(arq)
    ws = wb["Instrucoes"]; ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 120
    for r in ws.iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical="top"); r[0].font = Font(bold=True)
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for quem in ("Felipe", "Vini"):
        ws = wb[quem]
        hdr = {c.value: c.column_letter for c in ws[1]}
        n = ws.max_row
        for c in ws[1]:
            c.font = Font(bold=True)
            if c.value in PREENCHER:
                c.fill = amarelo
        for col, w_ in (("nome", 26), ("cidade", 18), ("link_satelite", 14), ("link_rua", 14), ("tipo_imovel", 28), ("observacao", 40)):
            ws.column_dimensions[hdr[col]].width = w_
        for i in range(2, n + 1):
            for col in ("link_satelite", "link_rua"):
                cel = ws[f"{hdr[col]}{i}"]
                cel.hyperlink = cel.value; cel.value = "abrir"; cel.font = Font(color="0563C1", underline="single")
            ws[f"{hdr['metragem_estimada_m2']}{i}"] = f'=IF(AND(ISNUMBER({hdr["area_telhado_m2"]}{i}),ISNUMBER({hdr["pavimentos_academia"]}{i})),{hdr["area_telhado_m2"]}{i}*{hdr["pavimentos_academia"]}{i},"")'
        for campo, opcoes in [("tipo_imovel", TIPOS), ("confianca_medida", CONF), ("mede_por_satelite", ["sim", "não"])] + \
                             [(c, SN) for c in ("letreiro_visivel_da_rua", "vitrine_para_rua", "frente_para_via_principal", "esquina", "estacionamento_proprio", "acesso_pelo_terreo")]:
            dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
            ws.add_data_validation(dv); dv.add(f"{hdr[campo]}2:{hdr[campo]}{n}")
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{hdr['nota_fachada_1a5']}2:{hdr['nota_fachada_1a5']}{n}")
        ws.freeze_panes = "C2"
    wb.save(arq)
    cont = base.groupby(["avaliador", "grupo"]).size().unstack(fill_value=0)
    print(cont.to_string()); print("prioridade (medir):", sem.prioridade.value_counts().to_dict())
    print("linhas por aba:", {q: int(base.avaliador.isin([q, "ambos"]).sum()) for q in ("Felipe", "Vini")})


if __name__ == "__main__":
    main()
