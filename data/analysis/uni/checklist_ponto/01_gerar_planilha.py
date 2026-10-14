# -*- coding: utf-8 -*-
"""BLK-UNI-17 (protocolo_checklist_ponto.md v1.0 revisada): planilhas do checklist do ponto, as cegas (sem faturamento,
alunos, ranking, grupo do estudo nem a hipotese). Avaliador A: obrigatorias (57 do modelo + turma de 2026 + turma B) em
ordem sorteada, opcionais no fim. Avaliador B: 20 sorteadas da 57 + turma 2026, com >= 4 de shopping, outra ordem.
Saida: planilha_checklist_ponto_A.xlsx, planilha_checklist_ponto_B.xlsx, grupos_unidades.csv (so' local) + copia das
duas planilhas no OneDrive (sem o protocolo). Rodar com PYTHONPATH=src (.venv)."""
import importlib.util
import shutil
import sys
import types
from pathlib import Path

try:
    import h3  # noqa: F401
except ImportError:
    for _k in [k for k in sys.modules if k == "h3" or k.startswith("h3.")]:
        del sys.modules[_k]
    sys.modules["h3"] = types.ModuleType("h3")

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

AQUI = Path(__file__).resolve().parent
UNI = AQUI.parent
OUT = Path(r"C:\Users\Felipe\OneDrive - Grupo Ultra\Área de Trabalho\Análises - Motor de Expansão\Exploratória BLK-UNI\10 - Checklist do ponto")
SEED = 20261007
CEGAS = ["belford-roxo-rj", "novo-gama-df", "paulinia-sp", "domingos-ferreira-pe", "bonsucesso-rj"]
VIS = ["sim", "não", "sem imagem utilizável"]
SNS = ["sim", "não", "não sei"]
OPCOES = {
    "imagem_antes_da_abertura": ["sim", "não"], "imovel_conferido": ["sim", "não"],
    "frente_visivel_da_via": ["sim", "não - obstruída", "sem imagem utilizável"], "fachada_envidracada": VIS, "esquina": VIS,
    "tipo_via": ["avenida", "rua", "interna"], "estacionamento": ["próprio do imóvel", "compartilhado", "só na rua", "sem imagem utilizável"],
    "entrada": ["térreo pela rua", "andar de cima com porta na rua", "por dentro de shopping ou galeria"],
    "supermercado_150m": SNS, "shopping_150m": SNS, "terminal_150m": SNS, "faculdade_150m": SNS, "quadra_comercial": SNS,
    "ponto_onibus_100m": SNS, "academia_concorrente_150m": SNS,
    "tipo_imovel": ["loja de rua", "galpão", "centro comercial ou prédio", "dentro de shopping", "anexo a supermercado"],
    "sabe_como_vai": ["sim", "não"],
}
ORDEM = ["imagem_antes_da_abertura", "data_imagem_rua", "imovel_conferido", "frente_visivel_da_via", "fachada_envidracada", "esquina", "testada_m",
         "tipo_via", "estacionamento", "entrada", "supermercado_150m", "shopping_150m", "terminal_150m", "faculdade_150m", "quadra_comercial",
         "ponto_onibus_100m", "academia_concorrente_150m", "tipo_imovel", "nota_potencial_1a5", "sabe_como_vai", "observacao"]
DEPOIS = ("nota_potencial_1a5", "sabe_como_vai", "observacao")
INSTRUCOES = [
    ("Objetivo", "Registrar características do imóvel de cada unidade pelo Street View e pelo satélite do Google Maps."),
    ("Regra 1 — só a imagem", "Use só a imagem da rua e o satélite. NÃO abra o cartão da academia no Google Maps (nota, avaliações, horários de pico, fotos) nem consulte números da unidade."),
    ("Regra 2 — sozinho", "Preencha sem ver a planilha de outra pessoa e sem trocar ideia sobre as unidades."),
    ("Regra 3 — o imóvel antes da academia", "No Street View, clique em 'Ver mais datas' e use a última imagem ANTERIOR à inauguração (coluna inaugurou). Se não houver, use a primeira depois e marque imagem_antes_da_abertura = não. Avalie o imóvel, não a academia pronta (letreiro, equipamentos)."),
    ("Como abrir", "Clique em 'abrir' na coluna link_rua. Se não abrir na frente do imóvel, use o link_satelite e arraste o bonequinho para a rua da entrada. A coluna confianca_coordenada avisa quando o ponto pode estar deslocado: confira e preencha imovel_conferido."),
    ("Ordem", "Preencha as linhas na ordem da planilha. As 'opcional' ficam no fim. Cronometre as 5 primeiras unidades e anote o tempo em observacao."),
    ("frente_visivel_da_via", "Da faixa de trânsito mais próxima, ~50 m antes do imóvel no sentido do trânsito, a frente aparece sem nada na frente (árvore, marquise, banca, outro prédio)? Se algo tapa: 'não - obstruída'. Sem imagem boa: 'sem imagem utilizável'."),
    ("fachada_envidracada", "A frente do imóvel no térreo tem vidro ou porta de vidro voltada para a calçada?"),
    ("esquina", "O imóvel fica no encontro de duas ruas que se cruzam? Lote que vai de uma rua até a paralela NÃO é esquina."),
    ("testada_m", "Largura da frente na rua da entrada, em metros: no satélite, botão direito > 'Medir distância'. Número inteiro. Dentro de shopping: deixe vazio."),
    ("tipo_via", "Conte as faixas de rolamento num sentido, sem a faixa de estacionamento. avenida = 2 ou mais faixas por sentido, ou canteiro central (mão única com 2 ou mais faixas conta); rua = 1 faixa por sentido; interna = dentro de shopping, condomínio ou estacionamento."),
    ("estacionamento", "próprio do imóvel; compartilhado (shopping, supermercado, centro comercial); só na rua; sem imagem utilizável."),
    ("entrada", "térreo pela rua; andar de cima com porta na rua; por dentro de shopping ou galeria."),
    ("Entorno (até 150 m)", "Medido a pé no Maps ('Medir distância'). NÃO conta o próprio imóvel nem o shopping ou supermercado onde ele fica. terminal = terminal ou estação de ônibus, metrô ou trem (ponto de ônibus comum não conta: ele vai em ponto_onibus_100m). quadra_comercial: a maioria dos imóveis no térreo da quadra é comércio? academia_concorrente_150m: academia de outra marca."),
    ("tipo_imovel", "loja de rua; galpão; centro comercial ou prédio; dentro de shopping; anexo a supermercado."),
    ("nota_potencial_1a5", "Só depois do checklist: o ponto, sozinho (sem pensar em cidade, renda ou concorrência), atrai aluno? 1 = muito pouco, 5 = muito. Número inteiro."),
    ("sabe_como_vai", "Você sabe se esta unidade vai bem ou mal? Responda com sinceridade: não é erro saber."),
    ("data_imagem_rua", "Data da imagem usada (aparece no canto do Street View), no formato aaaa-mm."),
]


def mod(nome, p):
    s = importlib.util.spec_from_file_location(nome, p); m = importlib.util.module_from_spec(s); s.loader.exec_module(m); return m


def planilha(base, arq, aba):
    with pd.ExcelWriter(arq, engine="openpyxl") as w:
        pd.DataFrame(INSTRUCOES, columns=["tema", "como fazer"]).to_excel(w, sheet_name="Instrucoes", index=False)
        base.to_excel(w, sheet_name=aba, index=False)
    wb = load_workbook(arq)
    wi = wb["Instrucoes"]; wi.column_dimensions["A"].width = 30; wi.column_dimensions["B"].width = 150
    ws = wb[aba]; hdr = {c.value: c.column_letter for c in ws[1]}; n = ws.max_row
    amarelo, azul = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="DDEBF7")
    for c in ws[1]:
        c.font = Font(bold=True)
        if c.value in ORDEM:
            c.fill = azul if c.value in DEPOIS else amarelo
    for col, w_ in (("unidade_id", 24), ("nome", 26), ("municipio", 18), ("link_rua", 10), ("link_satelite", 12), ("observacao", 40)):
        ws.column_dimensions[hdr[col]].width = w_
    for campo in ORDEM:
        if campo != "observacao":
            ws.column_dimensions[hdr[campo]].width = max(14, min(28, len(campo) + 2))
    for i in range(2, n + 1):
        for col in ("link_rua", "link_satelite"):
            cel = ws[f"{hdr[col]}{i}"]
            if cel.value:
                cel.hyperlink = cel.value; cel.value = "abrir"; cel.font = Font(color="0563C1", underline="single")
        ws[f"{hdr['data_imagem_rua']}{i}"].number_format = "@"
    def valida(dv, campo):
        dv.showErrorMessage = True; dv.error = "Use uma das opções da lista (ou deixe vazio)."; dv.errorTitle = "Valor fora da lista"
        ws.add_data_validation(dv); dv.add(f"{hdr[campo]}2:{hdr[campo]}{n}")
    for campo, ops in OPCOES.items():
        assert not any("," in o for o in ops), ops                     # virgula parte a lista do Excel
        valida(DataValidation(type="list", formula1='"' + ",".join(ops) + '"', allow_blank=True), campo)
    valida(DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True), "nota_potencial_1a5")
    valida(DataValidation(type="whole", operator="between", formula1="1", formula2="500", allow_blank=True), "testada_m")
    valida(DataValidation(type="textLength", operator="equal", formula1="7", allow_blank=True), "data_imagem_rua")
    ws.freeze_panes = "C2"
    wb.save(arq)


def main():
    u14 = mod("u14", UNI / "uni14" / "01_validacao_ampliada.py")
    Bs = u14.base(); d = Bs["d"]
    p1 = list(Bs["tr"].index) + list(Bs["te"].index)
    p2 = [u for u in pd.read_csv(UNI / "previsao_2026" / "previsoes_turma_2026.csv", sep=";", encoding="utf-8-sig", index_col=0).index if u not in p1]
    p3 = [u for u in u14.TURMA_B if u not in p1 + p2]
    ok = d.lat.notna() & d.lng.notna()
    p4 = [u for u in d.index if u not in p1 + p2 + p3 and ok.get(u, False)]
    assert set(CEGAS) <= set(p2) and not (set(p1) & set(p2) or set(p1 + p2) & set(p3))
    grupo = {**{u: "57_do_modelo" for u in p1}, **{u: ("cega_2026" if u in CEGAS else "turma_2026_exposta") for u in p2},
             **{u: "turma_B" for u in p3}, **{u: "demais" for u in p4}}
    rng = np.random.default_rng(SEED)
    obrig = list(rng.permutation(p1 + p2 + p3)); opc = list(rng.permutation(p4))
    base = pd.DataFrame({"unidade_id": obrig + opc, "preencher": ["obrigatória"] * len(obrig) + ["opcional"] * len(opc)})
    dd = d.reindex(base.unidade_id)
    base["nome"] = dd.nome.values
    base["municipio"] = dd.cres_nome_municipio_hex.fillna(dd.cidade_cadastro).values
    base["uf"] = dd.uf.values
    base["inaugurou"] = pd.to_datetime(dd.inauguracao.values).strftime("%Y-%m")
    base["confianca_coordenada"] = dd.confianca_coord.fillna("sem informação").values
    base["link_rua"] = [f"https://www.google.com/maps/@?api=1&map_action=pano&viewpoint={la:.6f},{lo:.6f}" if pd.notna(la) else ""
                        for la, lo in zip(dd.lat, dd.lng)]
    base["link_satelite"] = [f"https://www.google.com/maps/@{la:.6f},{lo:.6f},90m/data=!3m1!1e3" if pd.notna(la) else "" for la, lo in zip(dd.lat, dd.lng)]
    for c in ORDEM:
        base[c] = None
    sem_coord = list(base.unidade_id[base.link_rua == ""])
    shop = d.poi_em_shopping.reindex(p1 + p2).fillna(False).astype(bool)
    pool_s, pool_r = [u for u in p1 + p2 if shop[u] and ok.get(u, False)], [u for u in p1 + p2 if not shop[u] and ok.get(u, False)]
    lb = list(rng.choice(pool_s, size=min(4, len(pool_s)), replace=False)) + list(rng.choice(pool_r, size=20 - min(4, len(pool_s)), replace=False))
    B = base[base.unidade_id.isin(lb)].sample(frac=1, random_state=SEED + 1).reset_index(drop=True)
    B["preencher"] = "obrigatória"
    planilha(base, AQUI / "planilha_checklist_ponto_A.xlsx", "Avaliador_A")
    planilha(B, AQUI / "planilha_checklist_ponto_B.xlsx", "Avaliador_B")
    pd.DataFrame({"unidade_id": list(grupo), "grupo": list(grupo.values()), "na_planilha_B": [u in lb for u in grupo],
                  "em_shopping": [bool(d.poi_em_shopping.get(u)) if pd.notna(d.poi_em_shopping.get(u)) else None for u in grupo]}).to_csv(
        AQUI / "grupos_unidades.csv", sep=";", encoding="utf-8-sig", index=False)
    (AQUI / "lista_avaliador_B.csv").unlink(missing_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    for a in ("planilha_checklist_ponto_A.xlsx", "planilha_checklist_ponto_B.xlsx"):
        shutil.copy2(AQUI / a, OUT / a)
    print(f"A: 57 do modelo {len(p1)} | turma 2026 {len(p2)} (cegas {len(CEGAS)}) | turma B {len(p3)} | opcionais {len(p4)} | total {len(base)} | sem coordenada: {sem_coord}")
    print(f"B: {len(lb)} ({sum(shop.get(u, False) for u in lb)} de shopping): {sorted(str(u) for u in lb)}")
    print("coordenada media/baixa/sem info nas obrigatórias:", int((base.preencher.eq('obrigatória') & ~base.confianca_coordenada.eq('alta')).sum()))


if __name__ == "__main__":
    main()
