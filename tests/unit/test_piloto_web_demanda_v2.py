"""Viabilidade v2 nas superficies: rotas do piloto, planilha e PDF (DEC-078 ampliada).

Coeficientes FICTICIOS (os mesmos de `test_viabilidade_v2.py`), num `tmp_path`.
"""

from __future__ import annotations

import io
import json
import re
import sys
import warnings
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi import HTTPException

from motor_expansao.dimensionamento import viabilidade_v2 as v2
from motor_expansao.dimensionamento.payload_viabilidade import (
    ViabilidadeIn,
    investimento_do_body,
    montar_payload_para_pdf,
    premissas_do_body,
)

_REPO = Path(__file__).resolve().parents[2]
_SERVER = _REPO / "web" / "server"
if str(_SERVER) not in sys.path:
    sys.path.insert(0, str(_SERVER))

import acesso  # noqa: E402
import app as pilot  # noqa: E402

from tests.unit.dimensionamento.test_viabilidade_v2 import (  # noqa: E402
    COEF_FICTICIO,
    GOLDEN,
)


@pytest.fixture
def coef(tmp_path, monkeypatch):
    monkeypatch.delenv("MOTOR_VIABILIDADE_V2", raising=False)
    caminho = tmp_path / "v2_coeficientes.json"
    caminho.write_text(json.dumps(COEF_FICTICIO), encoding="utf-8")
    monkeypatch.setattr(v2, "caminho_coeficientes", lambda: caminho)
    v2.limpar_cache()
    yield v2.carregar_coeficientes()
    v2.limpar_cache()


@pytest.fixture
def sem_coef(tmp_path, monkeypatch):
    monkeypatch.delenv("MOTOR_VIABILIDADE_V2", raising=False)
    monkeypatch.setattr(v2, "caminho_coeficientes", lambda: tmp_path / "ausente.json")
    v2.limpar_cache()


def _sugestao(coef):
    corpo = v2.DemandaV2In(lat=GOLDEN["lat"], lng=GOLDEN["lng"], m2=GOLDEN["m2"], ticket=GOLDEN["ticket"])
    return v2.montar_sugestao_v2(corpo, pop_c10=3e5, dom_c10=1e5, renda_pc_c10=1800, coef=coef)


def _corpo_v2(sug, **kw):
    ref = {k: sug["entradas"][k] for k in ("pop_c10", "dom_c10", "renda_pc_c10")}
    return ViabilidadeIn(**{**GOLDEN, "demanda": round(sug["alunos"]["p50"]), "modelo_demanda": "v2",
                            "demanda_v2": {"token": sug["token"], **ref}, **kw})


# --------------------------------------------------------------------------- estado
def test_estado_sem_json_e_indisponivel_com_motivo(sem_coef):
    estado = pilot.demanda_v2_estado()
    assert estado["disponivel"] is False and "coeficientes" in estado["motivo"]


def test_estado_com_json_depende_da_rota_e_da_chave(coef, monkeypatch):
    monkeypatch.delenv("MOTOR_ROTAS_URL", raising=False)
    assert pilot.demanda_v2_estado()["disponivel"] is False
    monkeypatch.setenv("MOTOR_ROTAS_URL", "http://127.0.0.1:1")
    assert pilot.demanda_v2_estado()["disponivel"] is True
    monkeypatch.setenv("MOTOR_VIABILIDADE_V2", "0")
    estado = pilot.demanda_v2_estado()
    assert estado["disponivel"] is False and "desligada" in estado["motivo"]


# --------------------------------------------------------------------------- sugestao
_CORPO = {"lat": -23.33, "lng": -51.16, "m2": 1500}


def test_sugestao_sem_coeficientes_e_503(sem_coef, monkeypatch):
    monkeypatch.setenv("MOTOR_ROTAS_URL", "http://127.0.0.1:1")
    with pytest.raises(HTTPException) as exc:
        pilot.demanda_v2(v2.DemandaV2In(**_CORPO))
    assert exc.value.status_code == 503 and "coeficientes" in exc.value.detail


def test_sugestao_sem_rotas_e_503_e_nunca_cai_no_raio(coef, monkeypatch):
    monkeypatch.delenv("MOTOR_ROTAS_URL", raising=False)
    monkeypatch.setattr(pilot, "_entradas_c10_do_ponto", lambda *a: pytest.fail("não pode calcular"))
    with pytest.raises(HTTPException) as exc:
        pilot.demanda_v2(v2.DemandaV2In(**_CORPO))
    assert exc.value.status_code == 503 and exc.value.detail == v2.MOTIVO_SEM_ROTAS


def test_valhalla_fora_e_503_sem_ler_setor_nenhum(monkeypatch):
    from motor_expansao.api import service
    from motor_expansao.dashboard import alcance_ponto

    def _fora(*_a, **_k):
        raise alcance_ponto.RotasIndisponiveis("teste")

    monkeypatch.setattr(alcance_ponto, "isocrona_carro", _fora)
    monkeypatch.setattr(service, "_resolver_e_carregar", lambda *a, **k: pytest.fail("leu setor"))
    with pytest.raises(HTTPException) as exc:
        pilot._entradas_c10_do_ponto(-23.33, -51.16)
    assert exc.value.status_code == 503


def test_entradas_saem_da_isocrona_de_10_min_pela_geometria(monkeypatch):
    """A malha e' carregada PELA GEOMETRIA da isocrona (uniao das particoes), nunca pelo raio."""
    from shapely.geometry import box

    from motor_expansao.api import service
    from motor_expansao.dashboard import alcance_ponto
    from motor_expansao.dimensionamento import entradas_c10

    geom = box(-51.2, -23.4, -51.1, -23.3)
    chamadas = {}
    iso = SimpleNamespace(geometria=geom, area_km2=90.0, parametros={"minutos": 10})

    def _iso(lat, lng, minutos):
        chamadas["minutos"] = minutos
        return iso

    def _carregar(lat, lng, cfg, **kw):
        chamadas["kw"] = kw
        return "PR", "4113700", "SETORES"

    def _agregar(setores, g, lat, lng):
        assert setores == "SETORES" and g is geom
        return entradas_c10.EntradasC10(1e5, 3e4, 1500.0, 12, "calibrada")

    monkeypatch.setattr(alcance_ponto, "isocrona_carro", _iso)
    monkeypatch.setattr(service, "_resolver_e_carregar", _carregar)
    monkeypatch.setattr(entradas_c10, "agregar_entradas_c10", _agregar)
    entradas, recorte = pilot._entradas_c10_do_ponto(-23.33, -51.16)
    assert chamadas == {"minutos": 10, "kw": {"geometria": geom}}
    assert entradas.pop_c10 == 1e5 and recorte["minutos"] == 10 and recorte["n_setores"] == 12


def test_sugestao_devolve_faixa_alunos_flags_versao_e_token(coef, monkeypatch):
    from motor_expansao.dimensionamento.entradas_c10 import EntradasC10

    monkeypatch.setenv("MOTOR_ROTAS_URL", "http://127.0.0.1:1")
    monkeypatch.setattr(
        pilot, "_entradas_c10_do_ponto",
        lambda lat, lng: (EntradasC10(3e5, 1e5, 1800.0, 40, "calibrada"), {"minutos": 10}),
    )
    r = pilot.demanda_v2(v2.DemandaV2In(**_CORPO, ticket=147))
    assert set(r["faturamento"]) == {"p10", "p50", "p90"}
    assert r["alunos"]["p10"] < r["alunos"]["p50"] < r["alunos"]["p90"]
    assert r["faixa_provisoria"] is True and r["flag_extrapolacao"] is False
    assert r["coeficientes"]["sha256"] == coef.sha256 and r["token"].startswith("v2-")
    assert r["recorte"] == {"minutos": 10}
    assert r["entradas"]["renda_dom_c10"] == pytest.approx(5400.0)


# --------------------------------------------------------------------------- viabilidade
@pytest.fixture
def sem_catchment(monkeypatch):
    monkeypatch.setattr(pilot, "_setores_para_catchment", lambda *a: None)


def test_viabilidade_v2_token_velho_e_422(coef, sem_catchment):
    sug = _sugestao(coef)
    with pytest.raises(HTTPException) as exc:
        pilot._payload_viabilidade(_corpo_v2(sug, m2=1700))
    assert exc.value.status_code == 422 and "nova sugestão" in exc.value.detail


def test_viabilidade_v2_sem_coeficientes_e_503(coef, sem_catchment, monkeypatch, tmp_path):
    sug = _sugestao(coef)
    monkeypatch.setattr(v2, "caminho_coeficientes", lambda: tmp_path / "sumiu.json")
    with pytest.raises(HTTPException) as exc:
        pilot._payload_viabilidade(_corpo_v2(sug))
    assert exc.value.status_code == 503


def test_viabilidade_v2_ok(coef, sem_catchment):
    p = pilot._payload_viabilidade(_corpo_v2(_sugestao(coef)))
    assert p["demanda_fonte"] == "modelo_v2" and "demanda_v2" in p


# --------------------------------------------------------------------------- planilha
def _gerar_fake(capturado):
    def gerar(demanda, premissas, inv, *, rotulo=None, m2=None, aviso_nota=None,
              solicitante=None, report_id=None, demanda_v2=None):
        capturado["demanda_v2"] = demanda_v2
        return b"xlsx"

    return gerar


def test_rota_da_planilha_passa_o_bloco_so_no_modo_v2(coef, monkeypatch):
    capturado: dict = {}
    monkeypatch.setattr(pilot, "_gerador_simulador_xlsx", lambda: _gerar_fake(capturado))
    pilot._gerar_simulador_xlsx_response(ViabilidadeIn(**GOLDEN), "x")
    assert capturado["demanda_v2"] is None
    pilot._gerar_simulador_xlsx_response(_corpo_v2(_sugestao(coef)), "x")
    assert capturado["demanda_v2"]["demanda_fonte"] == "modelo_v2"
    with pytest.raises(HTTPException) as exc:
        pilot._gerar_simulador_xlsx_response(_corpo_v2(_sugestao(coef), m2=1600), "x")
    assert exc.value.status_code == 422


def _rotulos(wb) -> dict[str, tuple[int, object]]:
    ws = wb["Premissas"]
    return {str(ws.cell(row=r, column=1).value or ""): (r, ws.cell(row=r, column=2).value)
            for r in range(1, ws.max_row + 1)}


def test_planilha_v2_mostra_faturamento_fixo_e_r_aluno_como_formulas(coef, tmp_path):
    from motor_expansao.dimensionamento.simulador_xlsx import gerar_simulador_xlsx
    from motor_expansao.dimensionamento.viabilidade_v2 import conferir_demanda_v2

    body = _corpo_v2(_sugestao(coef))
    prem = premissas_do_body(body)
    bloco = conferir_demanda_v2(body, prem, rampa_reta_forcada=False).bloco
    xlsx = gerar_simulador_xlsx(float(body.demanda), prem, investimento_do_body(body), demanda_v2=bloco)
    wb = openpyxl.load_workbook(BytesIO(xlsx))
    rot = _rotulos(wb)
    assert rot["Faturamento maduro previsto (p50)"][1] == bloco["faturamento"]["p50"]
    assert rot["FIXO: faturamento com 0 alunos"][1].startswith("=Premissas!")
    assert rot["R_ALUNO: faturamento de cada aluno a mais"][1].startswith("=Premissas!")
    assert str(rot["Alunos que o faturamento p50 sustenta"][1]).startswith("=MAX(")
    assert "Sugerida pelo modelo de praça" in str(wb["Premissas"].cell(
        row=rot["Demanda total (alunos na maturidade)"][0], column=4).value)
    # Produto unico (emenda da DEC-078, 21/10/2026): nenhuma celula de TEXTO da planilha
    # fala em versao nem em faixa provisoria. Identificadores (nomes definidos `v2_*`) e
    # formulas nao sao texto de usuario.
    textos = [c.value for aba in wb.worksheets for linha in aba.iter_rows() for c in linha
              if isinstance(c.value, str) and not c.value.startswith("=")]
    assert [t for t in textos if re.search(r"\bv2\b|provis[óo]ri", t, re.I)] == []
    # Recalculo REAL (pacote `formulas`): as formulas da planilha batem com a conversao do motor.
    formulas = pytest.importorskip("formulas")
    caminho = tmp_path / "v2.xlsx"
    caminho.write_bytes(xlsx)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        sol = formulas.ExcelModel().loads(str(caminho)).finish().calculate()

    def ler(rotulo: str) -> float:
        valor = sol[f"'[v2.xlsx]PREMISSAS'!B{rot[rotulo][0]}"]
        return float(getattr(valor, "value", [[valor]])[0][0])

    assert ler("Alunos que o faturamento p50 sustenta") == pytest.approx(bloco["alunos"]["p50"], abs=0.06)
    assert ler("Alunos que o faturamento p10 sustenta") == pytest.approx(bloco["alunos"]["p10"], abs=0.06)
    assert ler("R_ALUNO: faturamento de cada aluno a mais") == pytest.approx(
        bloco["conversao"]["r_aluno"], rel=1e-6)
    assert ler("FIXO: faturamento com 0 alunos") == pytest.approx(bloco["conversao"]["fixo"], rel=1e-9)


def test_planilha_densidade_nao_tem_nada_da_v2():
    from motor_expansao.dimensionamento.simulador_xlsx import gerar_simulador_xlsx

    body = ViabilidadeIn(**GOLDEN)
    xlsx = gerar_simulador_xlsx(float(body.demanda), premissas_do_body(body), investimento_do_body(body))
    rot = _rotulos(openpyxl.load_workbook(BytesIO(xlsx)))
    assert not any("v2" in k or "R_ALUNO" in k for k in rot)


# --------------------------------------------------------------------------- PDF
def _texto_pdf(payload) -> str:
    import pypdf

    from motor_expansao.dashboard.censo_report import gerar_pdf_relatorio_pontual_classico

    resultado = {"lat": -23.33, "lng": -51.16, "nome_municipio": "LONDRINA", "uf": "PR", "raio_km": 1.0}
    pdf = gerar_pdf_relatorio_pontual_classico(resultado, None, viabilidade=payload)
    return "\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf)).pages)


def test_pdf_v2_carimba_a_origem_da_demanda(coef, tmp_path):
    payload = montar_payload_para_pdf(_corpo_v2(_sugestao(coef)), staging_dir=tmp_path)
    texto = _texto_pdf(payload)
    assert "Demanda: sugerida pelo modelo de praça." in " ".join(texto.split())
    assert "PREMISSA do operador" not in texto
    # produto unico (emenda de 21/10/2026): sem "v2" e sem aviso de faixa provisoria
    assert not re.search(r"\bv2\b", texto, re.I)
    assert "provisória" not in texto.lower()


def test_pdf_densidade_mantem_a_frase_de_sempre(tmp_path):
    texto = _texto_pdf(montar_payload_para_pdf(ViabilidadeIn(**GOLDEN), staging_dir=tmp_path))
    assert "PREMISSA do operador" in texto and "v2" not in texto


def test_linha_do_pdf_e_latin1_sem_pontuacao_tipografica():
    from motor_expansao.dashboard.censo_report import _viab_linha_demanda_v2

    assert _viab_linha_demanda_v2({}) is None
    dados = {
        "demanda_fonte": "premissa_editada",
        "demanda_v2": {"faturamento": {"p10": 1.0, "p50": 2.0, "p90": 3.0}, "demanda_sugerida": 2100,
                       "flag_extrapolacao": True, "rampa_reta_forcada_sem_agregador": True},
    }
    linha = _viab_linha_demanda_v2(dados)
    assert linha is not None
    assert "editada pelo operador (o modelo de praça sugeriu 2.100 alunos)" in linha
    assert "extrapolação" in linha
    # a reta nao e' mais forcada sem agregador; e nada de versao nem de "provisória"
    assert "reta" not in linha.lower()
    assert not re.search(r"\bv2\b|provis", linha, re.I)
    for aviso in ("nao_reconhecida", "indisponivel"):
        texto = _viab_linha_demanda_v2({"demanda_v2_aviso": aviso})
        assert texto is not None and not re.search(r"\bv2\b", texto, re.I)
        texto.encode("latin-1")
    linha.encode("latin-1")
    assert not any(c in linha for c in "—–•→…−“”")


# --------------------------------------------------------------------------- acesso
def test_rotas_da_v2_tem_o_gate_da_viabilidade():
    assert ("/api/demanda-v2", frozenset({"viabilidade"})) in acesso.REGRAS_DE_ACESSO
    assert acesso.superficie_necessaria("/api/demanda-v2") is None  # rota de malha
    assert acesso.rota_exige_malha_municipal("/api/demanda-v2/estado")
