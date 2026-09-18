"""As DUAS camadas de rastreio do simulador (D17, 10/09).

Planilha não tem content stream onde desenhar marca-d'água, então o carimbo é um par:
um bloco VISÍVEL na aba Aferição e as propriedades do documento, INVISÍVEIS, dentro do
ZIP. O que se guarda aqui é que as duas existem, que dizem a MESMA coisa, que sobrevivem
ao pós-processamento que reescreve o arquivo inteiro — e que sem identidade nenhuma o
arquivo sai como saía antes.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO

import openpyxl
import pytest

from motor_expansao.dashboard import pdf_base
from motor_expansao.dimensionamento.simulador import Premissas
from motor_expansao.dimensionamento.simulador_xlsx import (
    ABA_AFERICAO,
    gerar_simulador_xlsx,
)

RID = "f4ac2959-a584-4f86-9924-480302aeaff3"
QUEM = "will.lindo"


# Mesmo caso golden do resto do ciclo (Boulevard Londrina).
_DEMANDA = 2304.0
_INVEST = {"obra": 600_000.0, "equipamentos": 1_400_000.0, "taxa_franquia": 160_000.0}


def _gerar(**kw: object) -> bytes:
    premissas = Premissas(ticket_cheio=147.0, aluguel_mes=30_000.0, maturacao_meses=8)
    return gerar_simulador_xlsx(_DEMANDA, premissas, **_INVEST, **kw)  # type: ignore[arg-type]


@pytest.fixture(scope="module")
def blob() -> bytes:
    return _gerar(solicitante=QUEM, report_id=RID)


@pytest.fixture(scope="module")
def sem_carimbo() -> bytes:
    return _gerar()


# --------------------------------------------------------------------------------------
# A camada INVISÍVEL — docProps
# --------------------------------------------------------------------------------------


def test_o_report_id_esta_em_custom_xml_com_chave_NOMEADA(blob: bytes) -> None:
    """`custom.xml` é o que se lê por programa, sem regex em prosa."""
    custom = zipfile.ZipFile(BytesIO(blob)).read("docProps/custom.xml").decode("utf-8")
    achado = re.search(r'name="report_id".*?>([^<]+)<', custom)
    assert achado and achado.group(1) == RID
    assert re.search(r'name="solicitante".*?>([^<]+)<', custom).group(1) == QUEM


def test_core_xml_leva_o_nome_em_creator(blob: bytes) -> None:
    """`creator` é o campo que o Windows mostra em Propriedades > Detalhes > Autores:
    quem achar o arquivo lê isso sem abrir o Excel e sem acesso ao sistema."""
    core = zipfile.ZipFile(BytesIO(blob)).read("docProps/core.xml").decode("utf-8")
    assert "Will Lindo" in core
    assert f"report_id={RID}" in core


def test_o_subject_e_o_MESMO_texto_da_marca_dos_PDFs(blob: bytes) -> None:
    """Uma pessoa, um nome, nos quatro artefatos. É o mesmo anti-drift que pegou o
    `felipe_castaldi` saindo com underscore num PDF e sem noutro."""
    core = zipfile.ZipFile(BytesIO(blob)).read("docProps/core.xml").decode("utf-8")
    assert pdf_base.texto_da_marca(QUEM, RID) in core


def test_o_carimbo_sobrevive_a_reescrita_do_zip(blob: bytes) -> None:
    """`_com_valores_em_cache` reescreve TODO o ZIP para gravar o cache das fórmulas.
    Se ele descartasse `docProps`, a camada invisível morreria em silêncio — e o
    `gerar_simulador_xlsx` já devolve o arquivo depois dessa etapa."""
    partes = zipfile.ZipFile(BytesIO(blob)).namelist()
    assert "docProps/custom.xml" in partes
    assert "docProps/core.xml" in partes


def test_o_arquivo_reabre_no_openpyxl_com_as_propriedades(blob: bytes) -> None:
    wb = openpyxl.load_workbook(BytesIO(blob))
    assert wb.properties.creator == "Will Lindo"
    assert RID in (wb.properties.subject or "")


# --------------------------------------------------------------------------------------
# A camada VISÍVEL — o bloco na Aferição
# --------------------------------------------------------------------------------------


def _celulas(blob: bytes, aba: str) -> list[str]:
    ws = openpyxl.load_workbook(BytesIO(blob))[aba]
    return [str(c.value) for linha in ws.iter_rows() for c in linha if c.value is not None]


def test_o_bloco_visivel_esta_na_AFERICAO(blob: bytes) -> None:
    textos = _celulas(blob, ABA_AFERICAO)
    assert "RASTREIO DESTE ARQUIVO" in textos
    assert RID in textos
    assert "Will Lindo" in textos


def test_o_bloco_visivel_NAO_esta_no_Resumo_nem_nas_Premissas(blob: bytes) -> None:
    """De propósito: o simulador existe para ser aberto na frente do investidor, e essas
    são as abas que ele lê. A Aferição já é a aba de proveniência."""
    for aba in ("Resumo", "Premissas"):
        assert RID not in _celulas(blob, aba), f"vazou para a aba {aba}"


def test_o_bloco_visivel_diz_o_que_fazer_com_o_identificador(blob: bytes) -> None:
    """Um UUID sem instrução é ruído: quem abre não sabe que aquilo serve para alguma
    coisa, e a primeira reação a um número sem explicação é apagá-lo."""
    junto = " ".join(_celulas(blob, ABA_AFERICAO))
    assert "Estratégia e Growth" in junto
    assert "não é senha" in junto.lower() or "Não é senha" in junto


def test_as_duas_camadas_dizem_o_MESMO_id(blob: bytes) -> None:
    """Duas verdades sobre a mesma coisa é pior que uma: se divergissem, o rastreio
    daria respostas diferentes conforme onde se olhasse."""
    custom = zipfile.ZipFile(BytesIO(blob)).read("docProps/custom.xml").decode("utf-8")
    do_zip = re.search(r'name="report_id".*?>([^<]+)<', custom).group(1)
    assert do_zip in _celulas(blob, ABA_AFERICAO)


# --------------------------------------------------------------------------------------
# O ramo sem identidade
# --------------------------------------------------------------------------------------


def test_sem_identidade_nenhuma_o_arquivo_nao_ganha_bloco(sem_carimbo: bytes) -> None:
    """Chamada direta de biblioteca (teste, script) sai como saía. Mesmo contrato do
    `aviso_nota=None`."""
    assert "RASTREIO DESTE ARQUIVO" not in _celulas(sem_carimbo, ABA_AFERICAO)
    custom = zipfile.ZipFile(BytesIO(sem_carimbo)).namelist()
    assert "docProps/custom.xml" not in custom


def test_so_o_report_id_sem_pessoa_ainda_carimba() -> None:
    """Autoria nula é prevista (D19) e meio rastreio vale mais que nenhum: sem saber
    quem, ainda se sabe QUE geração produziu o arquivo."""
    blob = _gerar(report_id=RID)
    assert RID in _celulas(blob, ABA_AFERICAO)
    custom = zipfile.ZipFile(BytesIO(blob)).read("docProps/custom.xml").decode("utf-8")
    assert "solicitante" not in custom


def test_o_bloco_nao_desloca_a_tabela_de_afericao(blob: bytes, sem_carimbo: bytes) -> None:
    """Ele é ACRESCENTADO no fim. Se empurrasse as linhas, toda referência de fórmula
    que aponta para a Aferição passaria a apontar para a célula errada."""
    com = _celulas(blob, ABA_AFERICAO)
    sem = _celulas(sem_carimbo, ABA_AFERICAO)
    assert com[: len(sem)] == sem
