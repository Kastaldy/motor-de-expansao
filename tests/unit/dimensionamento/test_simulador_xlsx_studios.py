"""Simulador XLSX com STUDIOS de ticket proprio (pedido de 2026-10-05).

O motor passou a tratar cada studio (0..3) como uma FRACAO fixa da demanda total
(8% por studio) que paga o ticket DELE, com o tratamento do balcao (churn,
inadimplencia, reajuste anual e anuidade); o split balcao/agregador vale so' sobre a
demanda RESTANTE, e o custo fixo de cada studio entra nos outros fixos. A planilha
de formulas vivas tem de reproduzir isso ao centavo.

Dois niveis, como em `test_simulador_xlsx.py`:
1) ESTRUTURA (openpyxl): rotulo "Ticket de musculacao", 3 slots de ticket SEMPRE
   escritos (mesmo com 0 studio), validacao de lista no numero de studios, linhas
   novas na DRE e valores do motor na Afericao.
2) RECALCULO REAL (pacote `formulas`, SKIP sem ele): DRE e fluxo nos 64 meses, a
   Afericao sem nenhum DIVERGENTE e o Resumo contra os KPIs do motor, em DOIS
   cenarios com studio:
   - "dois": 1.000 alunos, 2 studios (157, 167), anuidade so' do balcao (default) ->
     80 + 80 nos studios, split sobre 840 (579,6 balcao / 260,4 agregador);
   - "tres": 2.304 alunos, 3 studios (157, 167, 177), anuidade cobrada de TODOS.
"""

from __future__ import annotations

import re
import unicodedata
import warnings
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from motor_expansao.dimensionamento.config import (
    SIM_CUSTO_STUDIO,
    SIM_TICKETS_STUDIO_PADRAO,
)
from motor_expansao.dimensionamento.simulador import Premissas, simular
from motor_expansao.dimensionamento.simulador_xlsx import (
    _DRE_ROW,
    _FLX_ROW,
    _MES_COL_INI,
    ABA_AFERICAO,
    ABA_DRE,
    ABA_FLUXO,
    ABA_PREMISSAS,
    ABA_RESUMO,
    gerar_simulador_xlsx,
)

_INVEST = {
    "obra": 600_000.0,
    "parcelas_obra": 4,
    "equipamentos": 1_400_000.0,
    "prazo_equipamentos": 60,
    "juros_equipamentos_am": 0.018,
    "taxa_franquia": 160_000.0,
}

# Unidade pequena para os 1.000 alunos do exemplo: com o investimento do golden o
# payback nao acontece no horizonte e a Afericao compararia "Nao atingido" com inf.
_INVEST_PEQUENO = {
    "obra": 200_000.0,
    "parcelas_obra": 4,
    "equipamentos": 250_000.0,
    "prazo_equipamentos": 36,
    "juros_equipamentos_am": 0.015,
    "taxa_franquia": 160_000.0,
}

_CENARIOS: dict[str, tuple[float, Premissas, dict[str, float]]] = {
    "dois": (
        1_000.0,
        Premissas(
            ticket_cheio=147.0, aluguel_mes=6_000.0, maturacao_meses=8,
            tickets_studios=(157.0, 167.0),
        ),
        _INVEST_PEQUENO,
    ),
    "tres": (
        2_304.0,
        Premissas(
            ticket_cheio=147.0, aluguel_mes=30_000.0, maturacao_meses=8,
            tickets_studios=(157.0, 167.0, 177.0), anuidade_apenas_balcao=False,
        ),
        _INVEST,
    ),
}

# linha da DRE -> campo da serie do motor (inclui as linhas novas dos studios).
_DRE_VS_MOTOR = {
    "alunos_total": "alunos_total",
    "alunos_balcao": "alunos_balcao",
    "alunos_agregadores": "alunos_agregadores",
    "alunos_studios": "alunos_studios",
    "rec_studios": "receita_studios",
    "faturamento": "faturamento_mensal",
    "rec_anuidade": "receita_anuidade",
    "deducoes": "deducoes",
    "receita_liquida": "receita_liquida",
    "impostos": "impostos",
    "receita_pos_impostos": "receita_pos_impostos",
    "cvar_total": "custos_variaveis",
    "folha": "folha",
    "outros_total": "outros_fixos",
    "of_studios": "custo_studios",
    "aluguel": "aluguel",
    "custo_pre_op": "custo_pre_operacional",
    "custos_op": "custos_op",
    "ebitda": "ebitda_mensal",
    "ir_total": "ir_csll",
    "juros": "juros",
}
_FLX_VS_MOTOR = {
    "ebitda": "ebitda_mensal",
    "ir_csll": "ir_csll",
    "juros": "juros",
    "pmt": "pmt",
    "amortizacao": "amortizacao",
    "investimento": "investimento",
    "fcf": "fcf_mensal",
    "fcf_acumulado": "fcf_acumulado",
}


def _norm(texto: object) -> str:
    sem = unicodedata.normalize("NFKD", str(texto))
    sem = "".join(c for c in sem if not unicodedata.combining(c))
    for traco in ("—", "–", "−"):
        sem = sem.replace(traco, "-")
    return re.sub(r"\s+", " ", sem).strip().lower()


def _premissas_por_rotulo(wb: openpyxl.Workbook) -> dict[str, tuple[int, object]]:
    ws = wb[ABA_PREMISSAS]
    return {
        _norm(ws.cell(row=r, column=1).value or ""): (r, ws.cell(row=r, column=2).value)
        for r in range(4, ws.max_row + 1)
    }


def _afericao(wb: openpyxl.Workbook) -> dict[str, tuple[int, object]]:
    ws = wb[ABA_AFERICAO]
    out: dict[str, tuple[int, object]] = {}
    for row in range(1, ws.max_row + 1):
        rotulo = ws.cell(row=row, column=1).value
        formula = ws.cell(row=row, column=3).value
        if rotulo and isinstance(formula, str) and formula.startswith("="):
            out[_norm(rotulo)] = (row, ws.cell(row=row, column=2).value)
    return out


@pytest.fixture(scope="module", params=sorted(_CENARIOS))
def cenario(request: pytest.FixtureRequest) -> dict[str, object]:
    demanda, p, invest = _CENARIOS[request.param]
    blob = gerar_simulador_xlsx(demanda, p, nome_ponto=f"Studios {request.param}", **invest)
    return {
        "nome": request.param,
        "demanda": demanda,
        "p": p,
        "blob": blob,
        "wb": openpyxl.load_workbook(BytesIO(blob), data_only=False),
        "r": simular(demanda, p, **invest),
    }


@pytest.fixture(scope="module")
def wb_sem_studio() -> openpyxl.Workbook:
    p = Premissas(ticket_cheio=147.0, aluguel_mes=30_000.0, maturacao_meses=8)
    return openpyxl.load_workbook(
        BytesIO(gerar_simulador_xlsx(2_304.0, p, **_INVEST)), data_only=False
    )


# ---------------------------------------------------------------------------
# Nivel 1 — estrutura
# ---------------------------------------------------------------------------


def test_rotulo_e_ticket_de_musculacao_e_nao_ticket_cheio(cenario) -> None:
    prem = _premissas_por_rotulo(cenario["wb"])
    assert prem["ticket de musculacao"][1] == 147.0
    for aba in (ABA_PREMISSAS, ABA_RESUMO, ABA_AFERICAO):
        ws = cenario["wb"][aba]
        for row in ws.iter_rows():
            for cel in row:
                if isinstance(cel.value, str) and not cel.value.startswith("="):
                    assert "ticket cheio" not in _norm(cel.value), (
                        f"{aba}!{cel.coordinate}: {cel.value!r}"
                    )


def test_tres_slots_de_ticket_mesmo_sem_studio(wb_sem_studio) -> None:
    """Estrutura FIXA: trocar 0 -> 2 studios dentro do Excel nao pode faltar linha."""
    prem = _premissas_por_rotulo(wb_sem_studio)
    assert prem["numero de studios"][1] == 0
    for i, padrao in enumerate(SIM_TICKETS_STUDIO_PADRAO, start=1):
        row, valor = prem[f"ticket do studio {i}"]
        assert valor == padrao
        nota = wb_sem_studio[ABA_PREMISSAS].cell(row=row, column=2).comment
        assert nota is not None and "INATIVO" in nota.text
    assert prem["fracao da demanda total atendida por studio"][1] == pytest.approx(0.08)
    assert prem["custo fixo por studio"][1] == SIM_CUSTO_STUDIO


def test_slots_ativos_levam_o_ticket_do_operador_sem_nota_de_inativo(cenario) -> None:
    p: Premissas = cenario["p"]
    prem = _premissas_por_rotulo(cenario["wb"])
    ws = cenario["wb"][ABA_PREMISSAS]
    assert prem["numero de studios"][1] == p.n_studios
    for i in range(1, 4):
        row, valor = prem[f"ticket do studio {i}"]
        if i <= p.n_studios:
            assert valor == p.tickets_studios[i - 1]
            assert ws.cell(row=row, column=2).comment is None
        else:
            assert valor == SIM_TICKETS_STUDIO_PADRAO[i - 1]
            assert "INATIVO" in ws.cell(row=row, column=2).comment.text


def test_numero_de_studios_tem_validacao_de_lista_0_a_3(cenario) -> None:
    ws = cenario["wb"][ABA_PREMISSAS]
    row = _premissas_por_rotulo(cenario["wb"])["numero de studios"][0]
    alvo = f"B{row}"
    achou = [
        dv for dv in ws.data_validations.dataValidation
        if dv.type == "list" and alvo in str(dv.sqref)
    ]
    assert achou, "numero de studios sem validacao de lista"
    assert achou[0].formula1 == '"0,1,2,3"'


def test_dre_tem_as_linhas_dos_studios_e_o_custo_dentro_dos_outros_fixos(cenario) -> None:
    ws = cenario["wb"][ABA_DRE]
    assert "studios" in _norm(ws.cell(row=_DRE_ROW["alunos_studios"], column=1).value)
    assert "studios" in _norm(ws.cell(row=_DRE_ROW["rec_studios"], column=1).value)
    assert "studios" in _norm(ws.cell(row=_DRE_ROW["of_studios"], column=1).value)
    # Ordem: alunos de studio e receita de studio entre agregadores e personal.
    assert _DRE_ROW["alunos_agregadores"] < _DRE_ROW["alunos_studios"] < _DRE_ROW["rec_balcao"]
    assert _DRE_ROW["rec_agregadores"] < _DRE_ROW["rec_studios"] < _DRE_ROW["rec_personal"]
    letra = get_column_letter(_MES_COL_INI + 10)
    total = str(ws[f"{letra}{_DRE_ROW['outros_total']}"].value)
    assert f"{letra}{_DRE_ROW['of_studios']}" in total, total
    assert f"{letra}{_DRE_ROW['rec_studios']}" in str(
        ws[f"{letra}{_DRE_ROW['faturamento']}"].value
    )


def test_afericao_grava_os_valores_do_motor_dos_studios(cenario) -> None:
    r = cenario["r"]
    p: Premissas = cenario["p"]
    afer = _afericao(cenario["wb"])
    esperado = {
        "ticket de musculacao": 147.0,
        "numero de studios": float(p.n_studios),
        "alunos de studios (steady)": r.alunos_studios_steady,
        "receita de mensalidades dos studios (steady)": r.receita_studios_mensal,
        "custo fixo dos studios (steady)": r.custo_studios_mensal,
        "ticket blended": r.ticket_blended,
        "custo fixo total sem aluguel (outros fixos + studios + folha)":
            p.custo_fixo_total_mes(cenario["demanda"]),
    }
    for rotulo, valor in esperado.items():
        assert rotulo in afer, f"rótulo {rotulo!r} ausente da Aferição"
        assert float(afer[rotulo][1]) == pytest.approx(float(valor), abs=0.01), rotulo
    assert r.alunos_studios_steady > 0 and r.receita_studios_mensal > 0
    assert r.custo_studios_mensal >= p.n_studios * SIM_CUSTO_STUDIO - 0.01


def test_sem_studio_as_linhas_novas_valem_zero_na_afericao(wb_sem_studio) -> None:
    afer = _afericao(wb_sem_studio)
    for rotulo in (
        "numero de studios", "alunos de studios (steady)",
        "receita de mensalidades dos studios (steady)", "custo fixo dos studios (steady)",
    ):
        assert afer[rotulo][1] == 0, rotulo


# ---------------------------------------------------------------------------
# Nivel 2 — recalculo real (pacote `formulas`)
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def ler(cenario, tmp_path_factory: pytest.TempPathFactory):
    formulas = pytest.importorskip(
        "formulas", reason="instale com `python -m pip install formulas` para recalcular"
    )
    caminho: Path = tmp_path_factory.mktemp(f"studios_{cenario['nome']}") / "sim.xlsx"
    caminho.write_bytes(cenario["blob"])
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        sol = formulas.ExcelModel().loads(str(caminho)).finish().calculate()
    base = f"'[{caminho.name}]"

    def _ler(aba: str, coord: str) -> float:
        chave = f"{base}{aba.upper()}'!{coord}"
        assert chave in sol, f"célula {aba}!{coord} não existe no modelo recalculado"
        valor = sol[chave]
        try:
            valor = valor.value[0, 0]
        except (AttributeError, TypeError, IndexError):
            pass
        assert not isinstance(valor, str), f"{aba}!{coord} devolveu texto/erro: {valor!r}"
        return float(valor)

    return _ler


def test_recalculo_reproduz_a_dre_e_o_fluxo_nos_64_meses(ler, cenario) -> None:
    divergencias: list[str] = []
    for j, linha in enumerate(cenario["r"].serie_mensal):
        letra = get_column_letter(_MES_COL_INI + j)
        for aba, mapa, rows in (
            (ABA_DRE, _DRE_VS_MOTOR, _DRE_ROW),
            (ABA_FLUXO, _FLX_VS_MOTOR, _FLX_ROW),
        ):
            for key, campo in mapa.items():
                obtido = ler(aba, f"{letra}{rows[key]}")
                esperado = float(linha[campo])
                if abs(obtido - esperado) > 0.01:
                    divergencias.append(
                        f"{aba}/{key} M{int(linha['mes'])}: planilha {obtido:.4f} "
                        f"vs motor {esperado:.4f}"
                    )
    assert not divergencias, "\n".join(divergencias[:30])


def test_recalculo_do_exemplo_do_pedido_80_80_e_split_sobre_840(ler, cenario) -> None:
    """1.000 alunos, 2 studios: 80 + 80 nos studios, 579,6 balcão e 260,4 agregador."""
    if cenario["nome"] != "dois":
        pytest.skip("o exemplo do pedido é o cenário de 2 studios")
    st = int(cenario["r"].mes_referencia_steady)
    letra = get_column_letter(_MES_COL_INI + 4 + st - 1)  # M-4..M-1 ocupam 4 colunas
    assert ler(ABA_DRE, f"{letra}{_DRE_ROW['alunos_total']}") == pytest.approx(1_000.0)
    assert ler(ABA_DRE, f"{letra}{_DRE_ROW['alunos_studios']}") == pytest.approx(160.0)
    assert ler(ABA_DRE, f"{letra}{_DRE_ROW['alunos_balcao']}") == pytest.approx(579.6)
    assert ler(ABA_DRE, f"{letra}{_DRE_ROW['alunos_agregadores']}") == pytest.approx(260.4)


def test_recalculo_da_afericao_nao_tem_nenhum_divergente(ler, cenario) -> None:
    ws = cenario["wb"][ABA_AFERICAO]
    conferidas = 0
    for row in range(1, ws.max_row + 1):
        rotulo = ws.cell(row=row, column=1).value
        motor = ws.cell(row=row, column=2).value
        formula = ws.cell(row=row, column=3).value
        if not (rotulo and isinstance(formula, str) and formula.startswith("=")):
            continue
        assert isinstance(motor, (int, float)), f"{rotulo}: motor não numérico ({motor!r})"
        obtido = ler(ABA_AFERICAO, f"C{row}")
        assert abs(obtido - float(motor)) < 0.01, f"{rotulo}: fórmula {obtido!r} vs {motor!r}"
        conferidas += 1
    assert conferidas >= 45


def test_recalculo_das_premissas_derivadas_bate_com_o_motor(ler, cenario) -> None:
    """ticket_blended / receita_por_aluno / fat_maduro / folha / custo fixo (com studios)."""
    p: Premissas = cenario["p"]
    demanda: float = cenario["demanda"]
    prem = _premissas_por_rotulo(cenario["wb"])

    def cel(rotulo: str) -> float:
        return ler(ABA_PREMISSAS, f"B{prem[rotulo][0]}")

    assert cel("ticket blended por aluno total") == pytest.approx(p.ticket_blended, abs=1e-6)
    assert cel("receita por aluno total (com anuidade)") == pytest.approx(
        p.receita_por_aluno_total, abs=1e-6
    )
    assert cel("faturamento maduro (base de dimensionamento da folha)") == pytest.approx(
        p.faturamento_maduro(demanda), abs=0.01
    )
    assert cel("folha mensal fixa (vale desde o mes 1)") == pytest.approx(
        p.folha_fixa_mes(demanda), abs=0.01
    )
    assert cel("custo fixo total, sem aluguel (outros fixos + studios + folha)") == (
        pytest.approx(p.custo_fixo_total_mes(demanda), abs=0.01)
    )
    assert cel("fracao da demanda total na musculacao (balcao + agregador)") == (
        pytest.approx(p.share_musculacao, abs=1e-12)
    )


def test_recalculo_do_resumo_bate_com_os_kpis_do_motor(ler, cenario) -> None:
    from motor_expansao.dimensionamento.simulador_xlsx import _RESUMO_ROW_INI, _linhas_resumo

    class _Refs(dict):
        def __missing__(self, chave: str) -> str:
            return f"{ABA_PREMISSAS}!$B$999"

    meses = [-4, -3, -2, -1, *range(1, 61)]
    linhas: dict[str, int] = {}
    row = _RESUMO_ROW_INI
    for key, *_resto in _linhas_resumo(meses, _Refs(), "BM", "BN"):
        if key:
            linhas[key] = row
        row += 1

    r = cenario["r"]
    esperado = {
        "faturamento": r.faturamento_mensal_steady,
        "receita_studios": r.receita_studios_mensal,
        "custo_studios": r.custo_studios_mensal,
        "folha": r.folha_mensal,
        "custos_op": r.custos_op_mensal,
        "ebitda": r.ebitda_mensal,
        "margem": r.margem_ebitda_pct,
        "break_even_ebitda": r.alunos_break_even_total,
        "break_even_caixa": r.alunos_break_even_caixa_total,
        "tir_anual": r.tir_anual,
        "vpl": r.vpl,
        "acumulado_m60": r.acumulado_mes_final,
        "ticket_blended": r.ticket_blended,
        "ticket_musculacao": 147.0,
        "n_studios": float(cenario["p"].n_studios),
        "teto_teto": r.aluguel_teto["teto"],
    }
    for key, valor in esperado.items():
        obtido = ler(ABA_RESUMO, f"B{linhas[key]}")
        tol = 0.0001 if abs(float(valor)) < 1 else 0.01
        assert obtido == pytest.approx(float(valor), abs=tol), key


# ---------------------------------------------------------------------------
# Modo "apenas balcao" trocado DENTRO do Excel: studios totalmente desligados
# ---------------------------------------------------------------------------


def test_modo_apenas_balcao_trocado_no_excel_desliga_receita_e_custo_dos_studios(
    tmp_path: Path,
) -> None:
    """O motor recusa studio no modo "apenas balcão", mas o operador pode trocar o modo
    na planilha. Ali os studios têm de sumir POR INTEIRO (alunos, receita, custo e as
    fatias derivadas), e a planilha tem de bater com o motor SEM studio nesse modo."""
    formulas = pytest.importorskip(
        "formulas", reason="instale com `python -m pip install formulas` para recalcular"
    )
    demanda, p, invest = _CENARIOS["dois"]
    assert p.n_studios == 2
    wb = openpyxl.load_workbook(BytesIO(gerar_simulador_xlsx(demanda, p, **invest)))
    prem = _premissas_por_rotulo(wb)
    row_modo = prem["modo da rampa"][0]
    assert prem["modo da rampa"][1] == "demanda total"
    wb[ABA_PREMISSAS].cell(row=row_modo, column=2).value = "apenas balcão"
    caminho = tmp_path / "balcao.xlsx"
    wb.save(caminho)

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        sol = formulas.ExcelModel().loads(str(caminho)).finish().calculate()
    base = f"'[{caminho.name}]"

    def ler(aba: str, coord: str) -> float:
        valor = sol[f"{base}{aba.upper()}'!{coord}"]
        try:
            valor = valor.value[0, 0]
        except (AttributeError, TypeError, IndexError):
            pass
        assert not isinstance(valor, str), f"{aba}!{coord} devolveu texto/erro: {valor!r}"
        return float(valor)

    def cel(rotulo: str) -> float:
        return ler(ABA_PREMISSAS, f"B{prem[rotulo][0]}")

    # Premissas: o número digitado continua 2, o EFETIVO vira 0 e zera o resto.
    assert cel("numero de studios") == 2
    assert cel('numero de studios efetivo (0 no modo "apenas balcao")') == 0
    assert cel("fracao da demanda total nos studios") == 0
    assert cel("custo fixo dos studios (total)") == 0
    assert cel("soma dos tickets dos studios ativos") == 0
    assert cel("receita de studio por aluno total") == 0

    # Referência: o motor no mesmo modo, sem studio.
    p0 = Premissas(
        ticket_cheio=p.ticket_cheio, aluguel_mes=p.aluguel_mes,
        maturacao_meses=p.maturacao_meses, rampa_apenas_balcao=True,
    )
    assert cel("ticket blended por aluno total") == pytest.approx(p0.ticket_blended, abs=1e-6)
    assert cel("receita por aluno total (com anuidade)") == pytest.approx(
        p0.receita_por_aluno_total, abs=1e-6
    )
    fat_maduro = cel("faturamento maduro (base de dimensionamento da folha)")
    assert fat_maduro == pytest.approx(
        demanda * cel("receita por aluno total (com anuidade)") + p0.personal_mes, abs=0.01
    )
    assert fat_maduro == pytest.approx(p0.faturamento_maduro(demanda), abs=0.01)
    assert cel("folha mensal fixa (vale desde o mes 1)") == pytest.approx(
        p0.folha_fixa_mes(demanda), abs=0.01
    )

    # DRE: linhas de studio zeradas nos 64 meses e o resto igual ao motor sem studio.
    r0 = simular(demanda, p0, **invest)
    divergencias: list[str] = []
    for j, linha in enumerate(r0.serie_mensal):
        letra = get_column_letter(_MES_COL_INI + j)
        for key in ("alunos_studios", "rec_studios", "of_studios"):
            assert ler(ABA_DRE, f"{letra}{_DRE_ROW[key]}") == 0, (key, j)
        for key, campo in _DRE_VS_MOTOR.items():
            obtido = ler(ABA_DRE, f"{letra}{_DRE_ROW[key]}")
            esperado = float(linha[campo])
            if abs(obtido - esperado) > 0.01:
                divergencias.append(
                    f"{key} M{int(linha['mes'])}: planilha {obtido:.4f} vs motor {esperado:.4f}"
                )
    assert not divergencias, "\n".join(divergencias[:30])
